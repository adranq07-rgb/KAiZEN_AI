from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from memory.learning_store import LearningStore


class AccountingCourseIngestor:
    def __init__(self, store=None, chunk_size=5000, domain="accounting"):
        self.store = store or LearningStore()
        self.chunk_size = chunk_size
        self.domain = domain

    @staticmethod
    def _subject(path):
        return path.stem.replace("-", " ").replace("_", " ")

    def _save_text(self, *, path, source_type, section, locator, text):
        chunks = []
        current = []
        size = 0
        for line in text.splitlines():
            if current and size + len(line) > self.chunk_size:
                chunks.append("\n".join(current))
                current, size = [], 0
            current.append(line)
            size += len(line) + 1
        if current:
            chunks.append("\n".join(current))
        saved = 0
        for index, chunk in enumerate(chunks, start=1):
            chunk_locator = f"{locator}; parca {index}/{len(chunks)}"
            if self.store.add(
                domain=self.domain,
                subject=self._subject(path),
                source_path=path,
                source_type=source_type,
                section=section,
                locator=chunk_locator,
                text=chunk,
            ):
                saved += 1
        return saved

    def ingest_xlsx(self, path):
        path = Path(path)
        workbook = load_workbook(path, read_only=True, data_only=False)
        saved = 0
        try:
            for sheet in workbook.worksheets:
                lines = []
                for row in sheet.iter_rows():
                    values = []
                    for cell in row:
                        if cell.value not in (None, ""):
                            values.append(f"{cell.coordinate}: {cell.value}")
                    if values:
                        lines.append(" | ".join(values))
                saved += self._save_text(
                    path=path,
                    source_type="xlsx_sheet",
                    section=sheet.title,
                    locator=f"sekme {sheet.title}",
                    text="\n".join(lines),
                )
        finally:
            workbook.close()
        return saved

    def ingest_pdf(self, path):
        path = Path(path)
        reader = PdfReader(str(path))
        saved = 0
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if len(text.strip()) < 40:
                continue
            saved += self._save_text(
                path=path,
                source_type="pdf_page",
                section=f"Sayfa {page_number}",
                locator=f"sayfa {page_number}",
                text=text,
            )
        return saved
