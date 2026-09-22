"""Kaizen'in kendi site toplayıcısı (Hermes'e bağımlı değil).

- robots.txt'e uyar, aynı siteye istekler arasında bekler, kendini tanıtır.
- sitemap.xml, RSS/Atom ve site içi bağlantılardan sayfa keşfeder.
- HTML, PDF ve Excel dosyalarının metnini Kaizen'in kendi deposuna yazar.
- Değişmeyen sayfayı tekrar indirmez (ETag / Last-Modified / son tarih).
- Forum, yorum ve arşiv-liste sayfalarını hafızaya almaz.
Sadece ağ hatasında (site erişilemezse) istenirse Hermes yedek okuyucu olarak devreye girer;
robots.txt veya 403/429 engelleri Hermes ile aşılmaya çalışılmaz, rapora yazılır.
"""
import gzip
import io
import re
import time
import urllib.error
import urllib.request
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from email.utils import parsedate_to_datetime
from urllib import robotparser
from urllib.parse import parse_qsl, quote, urlencode, urlparse, urlunparse

from accounting.html_extract import extract_html, normalize_date
from accounting.source_registry import path_allowed

USER_AGENT = "KaizenAI/1.0 (kisisel muhasebe arastirma asistani; robots.txt kurallarina uyar)"
ROBOT_TOKEN = "KaizenAI"

_SKIP_EXT = {
    ".jpg", ".jpeg", ".png", ".gif", ".webp", ".svg", ".ico", ".css", ".js", ".zip", ".rar", ".7z",
    ".mp3", ".mp4", ".avi", ".mov", ".woff", ".woff2", ".ttf", ".eot", ".exe", ".apk", ".doc", ".docx",
    ".ppt", ".pptx", ".xml", ".json", ".rss",
}
_GLOBAL_EXCLUDES = (
    "wp-admin", "wp-login", "wp-json", "xmlrpc", "cdn-cgi", "/login", "/giris", "/uye-", "/sepet", "/cart",
    "/checkout", "/yazdir", "/print", "/amp", "/feed", "?s=", "/search", "/arama", "/iletisim", "/gizlilik",
    "/kvkk", "/cerez",
)
_LISTING = re.compile(r"(?:/page/\d+|[?&]page=\d+|/kategori/|/category/|/tag/|/etiket/|/author/|/yazar/)", re.I)
_TRACKING = ("utm_", "fbclid", "gclid", "mc_", "ref", "replytocom")


@dataclass
class FetchResult:
    url: str
    status: int = 0
    content_type: str = ""
    body: bytes = b""
    headers: dict = field(default_factory=dict)
    final_url: str = ""
    error: str | None = None

    @property
    def ok(self):
        return self.status == 200 and self.error is None

    @property
    def blocked(self):
        return self.error == "robots" or self.status in (401, 403, 429)


def decode_body(body, content_type=""):
    match = re.search(r"charset=([\w-]+)", content_type or "", re.I)
    candidates = [match.group(1)] if match else []
    head = body[:3000].decode("ascii", "ignore")
    meta = re.search(r"charset=[\"']?([\w-]+)", head, re.I)
    if meta:
        candidates.append(meta.group(1))
    candidates += ["utf-8", "cp1254"]
    for enc in candidates:
        try:
            return body.decode(enc)
        except (LookupError, UnicodeDecodeError):
            continue
    return body.decode("utf-8", "replace")


class UrllibFetcher:
    def __init__(self, user_agent=USER_AGENT, timeout=25, min_delay=1.5, max_bytes=8_000_000,
                 respect_robots=True, sleep=time.sleep, clock=time.monotonic):
        self.user_agent = user_agent
        self.timeout = timeout
        self.min_delay = min_delay
        self.max_bytes = max_bytes
        self.respect_robots = respect_robots
        self._sleep, self._clock = sleep, clock
        self._last = {}
        self._robots = {}

    # --- düşük seviye -------------------------------------------------------------
    def _wait(self, host, delay):
        last = self._last.get(host)
        if last is not None:
            remaining = delay - (self._clock() - last)
            if remaining > 0:
                self._sleep(remaining)
        self._last[host] = self._clock()

    @staticmethod
    def _ascii_safe_url(url):
        """URL yolunda/sorgusunda Türkçe (ASCII dışı) karakter varsa yüzde-kodlar.

        Kodlanmamış Türkçe karakterli adresler (ör. .../Ağustos, .../yeni-yep'te-...) http.client'ın
        ascii kodlamasında patlıyordu ('ascii' codec can't encode character ...); zaten yüzde-kodlu
        kısımlar bozulmasın diye safe kümesine '%' eklendi (idempotent).
        """
        try:
            parsed = urlparse(url)
            path = quote(parsed.path, safe="/%:@!$&'()*+,;=~-")
            query = quote(parsed.query, safe="=&%:@!$'()*+,;~-")
            fragment = quote(parsed.fragment, safe="%:@!$&'()*+,;=~-")
            return urlunparse((parsed.scheme, parsed.netloc, path, parsed.params, query, fragment))
        except Exception:
            return url

    def _request(self, url, headers=None, delay=None):
        host = urlparse(url).netloc.lower()
        self._wait(host, self.min_delay if delay is None else delay)
        request = urllib.request.Request(self._ascii_safe_url(url), headers={
            "User-Agent": self.user_agent, "Accept-Encoding": "identity",
            "Accept": "text/html,application/xhtml+xml,application/xml,application/pdf,*/*;q=0.8",
            "Accept-Language": "tr-TR,tr;q=0.9,en;q=0.5", **(headers or {}),
        })
        try:
            with urllib.request.urlopen(request, timeout=self.timeout) as response:
                body = response.read(self.max_bytes + 1)
                if len(body) > self.max_bytes:
                    return FetchResult(url, response.status, error="too_large")
                return FetchResult(url, response.status, response.headers.get("Content-Type", ""), body,
                                   dict(response.headers), response.geturl())
        except urllib.error.HTTPError as exc:
            try:
                body = exc.read(200_000)
            except Exception:
                body = b""
            return FetchResult(url, exc.code, exc.headers.get("Content-Type", "") if exc.headers else "", body,
                               dict(exc.headers or {}), url)
        except (urllib.error.URLError, TimeoutError, OSError, ValueError) as exc:
            return FetchResult(url, 0, error=f"network: {exc}")

    # --- robots.txt ---------------------------------------------------------------
    def _robots_for(self, url):
        parsed = urlparse(url)
        base = f"{parsed.scheme}://{parsed.netloc}"
        if base in self._robots:
            return self._robots[base]
        result = self._request(base + "/robots.txt")
        parser = robotparser.RobotFileParser()
        entry = {"parser": None, "mode": "allow", "delay": None, "sitemaps": []}
        if result.status == 200:
            text = decode_body(result.body, result.content_type)
            parser.parse(text.splitlines())
            entry["parser"] = parser
            entry["delay"] = parser.crawl_delay(ROBOT_TOKEN) or parser.crawl_delay("*")
            entry["sitemaps"] = parser.site_maps() or []
        elif result.status in (401, 403):
            entry["mode"] = "deny"
        elif result.status == 404 or 400 <= result.status < 500:
            entry["mode"] = "allow"
        else:  # ağ hatası / 5xx: emin olamadık, bu koşuda dokunma
            entry["mode"] = "deny"
        self._robots[base] = entry
        return entry

    def robots_allows(self, url):
        entry = self._robots_for(url)
        if entry["mode"] == "deny":
            return False
        return True if entry["parser"] is None else entry["parser"].can_fetch(ROBOT_TOKEN, url)

    def sitemaps_for(self, url):
        return list(self._robots_for(url)["sitemaps"])

    def get(self, url, etag=None, last_modified=None):
        if self.respect_robots and not self.robots_allows(url):
            return FetchResult(url, 0, error="robots")
        headers = {}
        if etag:
            headers["If-None-Match"] = etag
        if last_modified:
            headers["If-Modified-Since"] = last_modified
        delay = self._robots_for(url)["delay"] if self.respect_robots else None
        return self._request(url, headers, delay=max(self.min_delay, float(delay)) if delay else None)


class HermesFallbackFetcher:
    """Site doğrudan açılamadığında (ağ hatası) Hermes'in web_extract aracıyla sayfayı okur."""

    def __init__(self, provider):
        self.provider = provider

    def get(self, url, **_):
        collect = getattr(self.provider, "collect_sources", None)
        if not callable(collect):
            return FetchResult(url, 0, error="fallback_unavailable")
        prompt = ("web_extract ile su sayfayi ac ve metnini oku: " + url + ". Yalniz bu araci kullan; "
                  "dosya veya ayar degistirme.")
        for item in collect(prompt) or []:
            if item.get("url", "").rstrip("/") == url.rstrip("/") and item.get("content"):
                return FetchResult(url, 200, "text/plain; charset=utf-8", item["content"].encode("utf-8"),
                                   {}, url)
        return FetchResult(url, 0, error="fallback_empty")


# --- sitemap / feed ayrıştırma -----------------------------------------------------

def _maybe_gunzip(body):
    if body[:2] == b"\x1f\x8b":
        try:
            return gzip.decompress(body)
        except OSError:
            return body
    return body


def _tag(block, name):
    match = re.search(rf"<(?:\w+:)?{name}[^>]*>(.*?)</(?:\w+:)?{name}>", block, re.S | re.I)
    if not match:
        return None
    value = match.group(1).strip()
    cdata = re.match(r"<!\[CDATA\[(.*?)\]\]>", value, re.S)
    return (cdata.group(1) if cdata else value).strip()


def parse_sitemap(xml_text):
    """Döner: (alt sitemap listesi [(url, lastmod)], sayfa listesi [(url, lastmod)])."""
    sitemaps = [(_tag(b, "loc"), _tag(b, "lastmod")) for b in re.findall(r"<(?:\w+:)?sitemap>(.*?)</(?:\w+:)?sitemap>", xml_text, re.S | re.I)]
    pages = [(_tag(b, "loc"), _tag(b, "lastmod")) for b in re.findall(r"<(?:\w+:)?url>(.*?)</(?:\w+:)?url>", xml_text, re.S | re.I)]
    return [(u, normalize_date(d)) for u, d in sitemaps if u], [(u, normalize_date(d)) for u, d in pages if u]


def parse_feed(xml_text):
    items = []
    for block in re.findall(r"<(?:item|entry)\b.*?</(?:item|entry)>", xml_text, re.S | re.I):
        link = _tag(block, "link")
        if not link:
            match = re.search(r"<link[^>]+href=[\"']([^\"']+)", block, re.I)
            link = match.group(1) if match else None
        stamp = _tag(block, "pubDate") or _tag(block, "published") or _tag(block, "updated")
        stamp_date = None
        if stamp:
            try:
                stamp_date = parsedate_to_datetime(stamp).date().isoformat()
            except (TypeError, ValueError):
                stamp_date = normalize_date(stamp)
        if link:
            items.append((link.strip(), stamp_date))
    return items


def extract_pdf_text(body, max_pages=80, max_chars=250_000):
    try:
        from pypdf import PdfReader
    except ImportError:
        return None, None
    try:
        reader = PdfReader(io.BytesIO(body))
        title = (getattr(reader.metadata, "title", None) or "") if reader.metadata else ""
        parts, total = [], 0
        for page in reader.pages[:max_pages]:
            text = (page.extract_text() or "").strip()
            if text:
                parts.append(text)
                total += len(text)
            if total >= max_chars:
                break
        return title, "\n\n".join(parts)[:max_chars]
    except Exception:
        return None, None


def extract_xlsx_text(body, max_rows=2000, max_chars=250_000):
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None
    try:
        workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception:
        return None
    lines = []
    try:
        for sheet in workbook.worksheets:
            lines.append(f"## Sekme: {sheet.title}")
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= max_rows:
                    break
                cells = [str(c) for c in row if c not in (None, "")]
                if cells:
                    lines.append(" | ".join(cells))
    finally:
        workbook.close()
    return "\n".join(lines)[:max_chars]


# --- Toplayıcı --------------------------------------------------------------------

def clean_url(url):
    parsed = urlparse(url.strip())
    query = [(k, v) for k, v in parse_qsl(parsed.query) if not k.lower().startswith(_TRACKING)]
    path = parsed.path or "/"
    return urlunparse((parsed.scheme.lower(), parsed.netloc.lower(), path, "", urlencode(query), ""))


def _host_matches(host, domains):
    host = (host or "").lower()
    return any(host == d or host.endswith("." + d) for d in domains)


class SiteCrawler:
    def __init__(self, store, fetcher=None, fallback=None, today=None, recheck_days=7):
        self.store = store
        self.fetcher = fetcher or UrllibFetcher()
        self.fallback = fallback
        self.today = today or date.today()
        self.recheck_days = recheck_days

    # --- kabul kuralları ----------------------------------------------------------
    def accept(self, source, url):
        try:
            parsed = urlparse(url)
        except ValueError:
            return False
        if parsed.scheme not in ("http", "https") or not _host_matches(parsed.hostname, source.domains):
            return False
        path = parsed.path.lower()
        if any(path.endswith(ext) for ext in _SKIP_EXT):
            return False
        lowered = (parsed.path + ("?" + parsed.query if parsed.query else "")).lower()
        if any(term in lowered for term in _GLOBAL_EXCLUDES):
            return False
        if any(term in path for term in source.excluded_path_terms):
            return False
        return path_allowed(source, url)

    @staticmethod
    def is_listing(url):
        return bool(_LISTING.search(url))

    # --- keşif ---------------------------------------------------------------------
    @staticmethod
    def bases(source):
        """Kaynağın kök adresleri: seed adreslerinin kökü, yoksa https://alan-adi."""
        origins = []
        for seed in source.seed_urls:
            parsed = urlparse(seed)
            origin = f"{parsed.scheme}://{parsed.netloc}"
            if origin not in origins:
                origins.append(origin)
        if not origins:
            origins = [f"https://{d}" for d in source.domains[:1]]
        return origins

    def _text(self, result):
        if not result.ok:
            return None
        return decode_body(_maybe_gunzip(result.body), result.content_type)

    def _discover_sitemaps(self, source, candidates, report, since, backfill, budget=25):
        seen, queue = set(), []
        for base in self.bases(source):
            if hasattr(self.fetcher, "sitemaps_for"):
                queue += [(u, None) for u in self.fetcher.sitemaps_for(base + "/")]
            queue += [(f"{base}/sitemap.xml", None), (f"{base}/sitemap_index.xml", None), (f"{base}/wp-sitemap.xml", None)]
        fetched = 0
        while queue and fetched < budget:
            url, _lastmod = queue.pop(0)
            if url in seen or not _host_matches(urlparse(url).hostname, source.domains):
                continue
            seen.add(url)
            result = self.fetcher.get(url)
            fetched += 1
            if result.blocked:
                report["blocked"].append(url)
                continue
            text = self._text(result)
            if not text or "<" not in text:
                continue
            children, pages = parse_sitemap(text)
            children.sort(key=lambda item: item[1] or "", reverse=True)
            for child, child_mod in children:
                if not backfill and child_mod and child_mod < since.isoformat():
                    continue
                queue.append((child, child_mod))
            for page_url, lastmod in pages:
                candidates.setdefault(clean_url(page_url), {"lastmod": lastmod, "via": "sitemap", "depth": 1})
            report["sitemaps"] += 1

    def _discover_feeds(self, source, candidates, report, feed_urls):
        base = self.bases(source)[0]
        for url in list(feed_urls) + [f"{base}/feed", f"{base}/rss.xml", f"{base}/feed/"]:
            result = self.fetcher.get(url)
            text = self._text(result)
            if not text or "<item" not in text.lower() and "<entry" not in text.lower():
                continue
            for link, stamp in parse_feed(text):
                candidates.setdefault(clean_url(link), {"lastmod": stamp, "via": "feed", "depth": 1})
            report["feeds"] += 1
            break

    # --- ana akış ------------------------------------------------------------------
    def crawl_source(self, source, max_pages=60, since_years=2, backfill=False, max_depth=2,
                     max_seconds=900, progress=None):
        """Bir kaynağı tarar; en yeni ve henüz bilinmeyen sayfalardan başlayarak max_pages sayfa okur."""
        started = time.monotonic()
        since = self.today - timedelta(days=365 * since_years)
        report = {
            "source": source.name, "mode": "crawl", "discovered": 0, "fetched": 0, "new": 0, "updated": 0,
            "unchanged": 0, "old_skipped": 0, "thin_skipped": 0, "errors": [], "blocked": [],
            "sitemaps": 0, "feeds": 0, "stopped": None,
        }
        if source.reference_only:
            report["mode"] = "reference_only"
            return report

        candidates = {}
        homes = list(source.seed_urls) or [b + "/" for b in self.bases(source)]
        for home in homes:
            candidates.setdefault(clean_url(home), {"lastmod": None, "via": "seed", "depth": 0})
        self._discover_sitemaps(source, candidates, report, since, backfill)
        if not report["sitemaps"]:
            self._discover_feeds(source, candidates, report, [])
        known = self.store.known_urls(source.name)

        def ordering(item):
            url, meta = item
            rank = {"seed": 0, "feed": 1, "sitemap": 1, "link": 2}[meta["via"]]
            lastmod = int(meta["lastmod"].replace("-", "")) if meta["lastmod"] else 0
            return (rank, -lastmod, url)

        visited, fetched_pages = set(), 0
        while fetched_pages < max_pages and time.monotonic() - started < max_seconds:
            pending = [(u, m) for u, m in candidates.items() if u not in visited and self.accept(source, u)]
            if not pending:
                break
            pending.sort(key=ordering)
            url, meta = pending[0]
            visited.add(url)
            info = known.get(url)

            if info and info["status"] in ("ok", "old") and meta["via"] != "seed":
                fetched_on = (info["fetched_at"] or "")[:10]
                if meta["lastmod"] and fetched_on and meta["lastmod"] <= fetched_on:
                    report["unchanged"] += 1
                    continue
                if not meta["lastmod"] and fetched_on and (
                    self.today - date.fromisoformat(fetched_on)
                ).days < self.recheck_days:
                    report["unchanged"] += 1
                    continue
            if not backfill and meta["lastmod"] and meta["lastmod"] < since.isoformat() and meta["via"] != "seed":
                report["old_skipped"] += 1
                continue

            result = self.fetcher.get(url, etag=(info or {}).get("etag"),
                                      last_modified=(info or {}).get("last_modified"))
            if result.error == "robots":
                report["blocked"].append(url)
                if meta["via"] == "seed" and not report["fetched"]:
                    report["stopped"] = "robots.txt bu siteye erişime izin vermiyor"
                    break
                continue
            if result.status == 0 and self.fallback is not None:
                result = self.fallback.get(url)
            if result.status == 304:
                report["unchanged"] += 1
                continue
            if result.status in (401, 403, 429):
                report["blocked"].append(url)
                if result.status == 429 or (meta["via"] == "seed" and not report["fetched"]):
                    report["stopped"] = f"site {result.status} döndürdü; bu koşuda durduruldu"
                    break
                continue
            if not result.ok:
                report["errors"].append(f"{url}: {result.error or result.status}")
                if len(report["errors"]) >= 15 and not report["fetched"]:
                    report["stopped"] = "art arda hata"
                    break
                continue

            fetched_pages += 1
            report["fetched"] += 1
            outcome, links = self._store_result(source, url, meta, result, since, backfill, report)
            if outcome:
                report[outcome] += 1
            if meta["depth"] < max_depth:
                for link in links:
                    link = clean_url(link)
                    if link not in candidates and self.accept(source, link):
                        candidates[link] = {"lastmod": None, "via": "link", "depth": meta["depth"] + 1}
            if progress:
                progress(source.name, fetched_pages, url)

        report["discovered"] = len(candidates)
        report["seconds"] = round(time.monotonic() - started, 1)
        self.store.set_state(source.name, "last_run", datetime.now().isoformat(timespec="seconds"))
        return report

    def _store_result(self, source, url, meta, result, since, backfill, report):
        """Sonucu depoya yazar. Döner: (yeni|updated|unchanged|None, bulunan bağlantılar)."""
        ctype = (result.content_type or "").lower()
        final_url = clean_url(result.final_url or url)
        kind, title, text, links = "html", "", "", []
        published = meta.get("lastmod")
        etag = result.headers.get("ETag") or result.headers.get("Etag")
        last_modified = result.headers.get("Last-Modified")
        common = {"via": meta["via"], "lastmod": meta.get("lastmod"), "crawler": "kaizen"}

        if "pdf" in ctype or url.lower().endswith(".pdf"):
            kind = "pdf"
            pdf_title, text = extract_pdf_text(result.body)
            title = pdf_title or url.rsplit("/", 1)[-1]
            if text is None:
                return self.store.upsert_page(url=final_url, source_name=source.name, title=title, kind="pdf",
                                              link_only=True, metadata=common), links
        elif "spreadsheet" in ctype or "excel" in ctype or re.search(r"\.xlsx?$", url.lower()):
            kind = "xlsx"
            text = extract_xlsx_text(result.body)
            title = url.rsplit("/", 1)[-1]
            if text is None:
                return self.store.upsert_page(url=final_url, source_name=source.name, title=title, kind="xlsx",
                                              link_only=True, metadata=common), links
        elif "text/plain" in ctype:
            text, title = decode_body(result.body, result.content_type), url
        else:
            page = extract_html(decode_body(result.body, result.content_type), final_url)
            title, text, links = page["title"], page["text"], page["links"]
            published = page["published_at"] or published or normalize_date(last_modified or "")

        if len(text) < 200:
            report["thin_skipped"] += 1
            return None, links
        if kind == "html" and self.is_listing(url):
            return None, links              # arşiv/liste sayfası: yalnız bağlantı keşfi için okundu
        if not backfill and published and published < since.isoformat():
            self.store.upsert_page(url=final_url, source_name=source.name, title=title, content="",
                                   kind=kind, published_at=published, status="old", metadata=common)
            report["old_skipped"] += 1
            return None, links
        outcome = self.store.upsert_page(
            url=final_url, source_name=source.name, title=title, content=text, kind=kind,
            published_at=published, etag=etag, last_modified=last_modified, metadata=common,
        )
        return outcome, links
